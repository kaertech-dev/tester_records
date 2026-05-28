"""
excel_form.py
Builds the Kaertech downtime report xlsx in memory.
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ──────────────────────────────────────────────
# Style helpers
# ──────────────────────────────────────────────

def thin_side():   return Side(style='thin',   color='000000')
def dash_side():   return Side(style='dashed', color='000000')
def none_side():   return Side(style='none')
def medium_side(): return Side(style='medium', color='000000')

def _font(bold=False, size=10, italic=False, name='Arial'):
    return Font(bold=bold, size=size, italic=italic, name=name)

def _align(h='left', v='center', wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _apply_border_box(ws, min_row, max_row, min_col, max_col, side_fn):
    s = side_fn()
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            left   = s if c == min_col else none_side()
            right  = s if c == max_col else none_side()
            top    = s if r == min_row else none_side()
            bottom = s if r == max_row else none_side()
            existing = cell.border
            cell.border = Border(
                left   = left   if c == min_col else existing.left,
                right  = right  if c == max_col else existing.right,
                top    = top    if r == min_row else existing.top,
                bottom = bottom if r == max_row else existing.bottom,
            )

def _fill_box(ws, min_row, max_row, min_col, max_col, side_fn):
    s = side_fn()
    b = Border(left=s, right=s, top=s, bottom=s)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = b


# ──────────────────────────────────────────────
# Image embedding helper
# ──────────────────────────────────────────────

# Column width is 2.2 units; 1 unit ≈ 7 px at 96 dpi → each col ≈ 15 px
_COL_WIDTH_UNITS = 2.2
_PX_PER_UNIT     = 7
_COL_PX          = int(_COL_WIDTH_UNITS * _PX_PER_UNIT)  # ≈ 15 px
_INNER_COLS      = 44   # columns C(3)…AT(46)
_TOTAL_ZONE_PX   = _INNER_COLS * _COL_PX                 # ≈ 660 px


def _embed_images_row(ws, img_bytes_list: list, anchor_row: int,
                      start_col: int = 3, zone_px_tall: int = 120):
    """
    Embed a list of images side-by-side.

    Each image is anchored to a cell string (e.g. "C17") — the simplest
    and most reliable openpyxl approach.  We calculate which column each
    image should start at based on how wide the previous images are,
    then just set xl_img.anchor = "<col><row>".
    """
    from openpyxl.drawing.image import Image as XLImage

    try:
        from PIL import Image as PILImage
        _has_pil = True
    except ImportError:
        _has_pil = False

    n = len(img_bytes_list)
    if n == 0:
        return

    gap_cols  = 1                                              # blank column between images
    slot_px   = max(30, (_TOTAL_ZONE_PX - (n - 1) * _COL_PX * gap_cols) // n)
    max_h_px  = max(10, zone_px_tall - 4)

    current_col = start_col   # 1-based column index, advances after each image

    for i, raw_bytes in enumerate(img_bytes_list):
        if not raw_bytes:
            continue
        try:
            stream = io.BytesIO(raw_bytes)

            if _has_pil:
                with PILImage.open(io.BytesIO(raw_bytes)) as pil:
                    nat_w, nat_h = pil.size
                if nat_w <= 0 or nat_h <= 0:
                    nat_w, nat_h = slot_px, max_h_px
                scale   = min(slot_px / nat_w, max_h_px / nat_h, 1.0)
                final_w = max(1, int(nat_w * scale))
                final_h = max(1, int(nat_h * scale))
            else:
                final_w = slot_px
                final_h = max_h_px

            xl_img        = XLImage(stream)
            xl_img.width  = final_w
            xl_img.height = final_h
            # Anchor: simple cell string — most reliable with openpyxl
            xl_img.anchor = f"{get_column_letter(current_col)}{anchor_row}"
            ws.add_image(xl_img)

            # Advance column cursor: how many cols does this image span?
            cols_used    = max(1, (final_w + _COL_PX - 1) // _COL_PX)
            current_col += cols_used + gap_cols

        except Exception as exc:
            print(f"[excel_form] image {i} embed error: {exc}")
            continue


# ──────────────────────────────────────────────
# Main public function
# ──────────────────────────────────────────────

def build_downtime_report(
    data: dict,
    # Legacy single-image kwargs (backwards-compat)
    analysis_img: bytes | None = None,
    verification_img: bytes | None = None,
    # Multi-image kwargs
    analysis_imgs: list | None = None,
    corrective_action_imgs: list | None = None,
    verification_imgs: list | None = None,
) -> bytes:

    # Merge legacy single-image into lists
    if analysis_imgs is None:
        analysis_imgs = [analysis_img] if analysis_img else []
    if corrective_action_imgs is None:
        corrective_action_imgs = []
    if verification_imgs is None:
        verification_imgs = [verification_img] if verification_img else []

    # Filter out None / empty
    analysis_imgs          = [b for b in analysis_imgs          if b]
    corrective_action_imgs = [b for b in corrective_action_imgs if b]
    verification_imgs      = [b for b in verification_imgs      if b]

    wb = Workbook()
    ws = wb.active
    ws.title = "Downtime Report"
    ws.sheet_view.showGridLines = False

    C_B  = 2
    C_AT = 46

    # ── Column widths ─────────────────────────────────────────────────────────
    for c in range(1, 47):
        ws.column_dimensions[get_column_letter(c)].width = _COL_WIDTH_UNITS

    # ── Row heights ───────────────────────────────────────────────────────────
    ROW_H = {
        1: 14, 2: 30, 3: 10, 4: 10, 5: 5,
        6: 16, 7: 16, 8: 5, 9: 16, 10: 5,
        11: 16, 12: 14, 13: 40, 14: 5,
        15: 16, 16: 14,
    }
    for r in range(17, 25): ROW_H[r] = 18
    ROW_H[24] = 5
    ROW_H[25] = 16
    ROW_H[26] = 14
    for r in range(27, 43): ROW_H[r] = 18
    ROW_H[42] = 5
    ROW_H[43] = 16
    ROW_H[44] = 14
    for r in range(45, 54): ROW_H[r] = 18
    ROW_H[53] = 5
    ROW_H[54] = 14
    ROW_H[55] = 14
    ROW_H[56] = 14
    ROW_H[57] = 14
    ROW_H[58] = 10

    for r, h in ROW_H.items():
        ws.row_dimensions[r].height = h

    # ── Outer border ──────────────────────────────────────────────────────────
    _apply_border_box(ws, 6, 57, C_B, C_AT, thin_side)

    # ── Row 1 – Company name ──────────────────────────────────────────────────
    ws.merge_cells("B1:AT1")
    ws["B1"] = "Kaertech Electronics Philippines Inc."
    ws["B1"].font      = _font(italic=True, size=9)
    ws["B1"].alignment = _align('center')

    # ── Row 2-4 – Title ───────────────────────────────────────────────────────
    ws.merge_cells("B2:AT4")
    ws["B2"] = "TROUBLESHOOTING / DOWNTIME REPORT"
    ws["B2"].font      = _font(bold=True, size=16)
    ws["B2"].alignment = _align('center', 'center')

    # ── EQUIPMENT (rows 6-9) ──────────────────────────────────────────────────
    ws.merge_cells("C6:AT6")
    ws["C6"] = "Equipment"
    ws["C6"].font      = _font(bold=True, size=10)
    ws["C6"].alignment = _align('left', 'center')
    for c in range(C_B, C_AT + 1):
        cell = ws.cell(row=6, column=c)
        cell.border = Border(left=cell.border.left, right=cell.border.right,
                             top=cell.border.top, bottom=dash_side())

    ws.merge_cells("C7:L7");   ws["C7"] = "Machine Brand/Model"
    ws["C7"].font = _font(size=9); ws["C7"].alignment = _align('left')
    ws.merge_cells("M7:U7");   ws["M7"] = data.get('asset_name', '')
    ws["M7"].alignment = _align('left')
    ws.merge_cells("V7:AE7");  ws["V7"] = "Machine Serial No."
    ws["V7"].font = _font(size=9); ws["V7"].alignment = _align('left')
    ws.merge_cells("AF7:AT7"); ws["AF7"] = data.get('asset_id', '')
    ws["AF7"].alignment = _align('left')
    for c in range(C_B, C_AT + 1):
        cell = ws.cell(row=7, column=c)
        cell.border = Border(left=cell.border.left, right=cell.border.right,
                             top=cell.border.top, bottom=dash_side())

    ws.merge_cells("C9:H9");  ws["C9"] = "Line No."
    ws["C9"].font = _font(size=9); ws["C9"].alignment = _align('left')
    ws.merge_cells("I9:AT9"); ws["I9"] = data.get('line_no', '')
    ws["I9"].alignment = _align('left')
    for c in range(C_B, C_AT + 1):
        cell = ws.cell(row=9, column=c)
        cell.border = Border(left=cell.border.left, right=cell.border.right,
                             top=cell.border.top, bottom=dash_side())

    # ── PROBLEM (rows 11-13) ──────────────────────────────────────────────────
    _apply_border_box(ws, 11, 13, C_B, C_AT, dash_side)
    ws.merge_cells("C11:AT11"); ws["C11"] = "Problem"
    ws["C11"].font = _font(bold=True, size=10); ws["C11"].alignment = _align('left')
    ws.merge_cells("C12:AT12"); ws["C12"] = "Description"
    ws["C12"].font = _font(size=9); ws["C12"].alignment = _align('left')
    ws.merge_cells("C13:AS13"); ws["C13"] = data.get('description', '')
    ws["C13"].alignment = _align('left', 'top')

    # ── ANALYSIS (rows 15-23) ─────────────────────────────────────────────────
    _apply_border_box(ws, 15, 23, C_B, C_AT, dash_side)
    ws.merge_cells("C15:AT15"); ws["C15"] = "Analysis"
    ws["C15"].font = _font(bold=True, size=10); ws["C15"].alignment = _align('left')
    ws.merge_cells("C16:AT16"); ws["C16"] = data.get('analysis', '')
    ws["C16"].alignment = _align('left', 'top', wrap=True)

    # rows 17-23 = 7 rows × 18pt × 1.33 px/pt ≈ 167 px tall
    if analysis_imgs:
        _embed_images_row(ws, analysis_imgs, anchor_row=17,
                          start_col=3, zone_px_tall=167)
    else:
        ws.merge_cells("C17:AS23")
        ws["C17"].alignment = _align('left', 'top')

    # ── CORRECTIVE ACTION (rows 25-41) ────────────────────────────────────────
    _apply_border_box(ws, 25, 41, C_B, C_AT, dash_side)
    ws.merge_cells("C25:AT25"); ws["C25"] = "Corrective Action"
    ws["C25"].font = _font(bold=True, size=10); ws["C25"].alignment = _align('left')
    ws.merge_cells("C26:AS40"); ws["C26"] = data.get('corrective_action', '')
    ws["C26"].alignment = _align('left', 'top', wrap=True)

    # rows 27-41 = 15 rows × 18pt × 1.33 ≈ 359 px tall
    if corrective_action_imgs:
        _embed_images_row(ws, corrective_action_imgs, anchor_row=27,
                          start_col=3, zone_px_tall=359)
    else:
        ws.merge_cells("C27:AS41")
        ws["C27"].alignment = _align('left', 'top')

    # ── VERIFICATION RESULT (rows 43-52) ──────────────────────────────────────
    _apply_border_box(ws, 43, 52, C_B, C_AT, dash_side)
    ws.merge_cells("C43:AT43"); ws["C43"] = "Verification Result:"
    ws["C43"].font = _font(bold=True, size=10); ws["C43"].alignment = _align('left')
    ws.merge_cells("C44:AT44"); ws["C44"] = data.get('verification_result', '')
    ws["C44"].alignment = _align('left', 'top', wrap=True)

    # rows 45-52 = 8 rows × 18pt × 1.33 ≈ 192 px tall
    if verification_imgs:
        _embed_images_row(ws, verification_imgs, anchor_row=45,
                          start_col=3, zone_px_tall=192)
    else:
        ws.merge_cells("C45:AS52")
        ws["C45"].alignment = _align('left', 'top')

    # ── DATE / TIME / PERSONNEL (rows 54-57) ──────────────────────────────────
    headers   = ["Date", "Time", "Troubleshoot by:"]
    col_spans = [(2, 14), (15, 24), (25, 46)]
    for (start, end), hdr in zip(col_spans, headers):
        ws.merge_cells(f"{get_column_letter(start)}54:{get_column_letter(end)}54")
        cell = ws.cell(row=54, column=start)
        cell.value     = hdr
        cell.font      = _font(bold=True, size=9)
        cell.alignment = _align('center')

    _apply_border_box(ws, 54, 57, C_B, C_AT, dash_side)

    def _split_dt(raw: str):
        if not raw:
            return '', ''
        try:
            dt = datetime.fromisoformat(raw.replace('T', ' '))
            return dt.strftime('%Y-%m-%d'), dt.strftime('%H:%M')
        except ValueError:
            parts = raw.split('T') if 'T' in raw else raw.split(' ')
            return (parts[0] if parts else raw, parts[1] if len(parts) > 1 else '')

    equip_down_date,   equip_down_time   = _split_dt(data.get('equip_down', ''))
    repair_start_date, repair_start_time = _split_dt(data.get('repair_start', ''))
    repair_end_date,   repair_end_time   = _split_dt(data.get('repair_end', ''))

    rows_data = [
        (55, "Equipt Down:",  equip_down_date,   equip_down_time,   data.get('troubleshoot_by', '')),
        (56, "Repair Start:", repair_start_date, repair_start_time, "Checked by:"),
        (57, "Repair End:",   repair_end_date,   repair_end_time,   "Approved by:"),
    ]

    for (row, lbl_date, date_val, time_val, lbl_person) in rows_data:
        ws.merge_cells(f"B{row}:G{row}");  ws[f"B{row}"] = lbl_date
        ws[f"B{row}"].font = _font(bold=True, size=9)
        ws[f"B{row}"].alignment = _align('left')
        ws.merge_cells(f"H{row}:N{row}");  ws[f"H{row}"] = date_val
        ws[f"H{row}"].alignment = _align('left')
        ws.merge_cells(f"O{row}:X{row}");  ws[f"O{row}"] = time_val
        ws[f"O{row}"].alignment = _align('left')
        ws.merge_cells(f"Y{row}:AT{row}"); ws[f"Y{row}"] = lbl_person
        ws[f"Y{row}"].font = _font(bold=(lbl_person not in ('', None)), size=9)
        ws[f"Y{row}"].alignment = _align('left')

    for row in range(55, 58):
        for c in range(C_B, C_AT + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = Border(left=cell.border.left, right=cell.border.right,
                                 top=dash_side(), bottom=dash_side())

    # ── Footer (row 58) ───────────────────────────────────────────────────────
    retention = data.get('retention_period', '5 years')
    eff_date  = data.get('effective_date', 'June 19, 2020')
    ws.merge_cells("B58:V58")
    ws["B58"] = f"Retention Period: {retention}          Effective Date: {eff_date}"
    ws["B58"].font = _font(italic=True, size=7)
    ws["B58"].alignment = _align('left', 'center')
    ws.merge_cells("W58:AT58")
    ws["W58"] = "FM-SMT-060 Rev.01"
    ws["W58"].font = _font(italic=True, size=7)
    ws["W58"].alignment = _align('right', 'center')

    # ── Restore outer thin border ─────────────────────────────────────────────
    _apply_border_box(ws, 6, 57, C_B, C_AT, thin_side)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()